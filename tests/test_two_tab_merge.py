"""QA de unir las filas de los dos tabs del invoice file (sin browser).

Desde el 2026-09-03 las filas del invoice file estan repartidas en dos pestañas
del mismo reporte: 'Vendor Entry Status', que lleva un filtro fijo
`Status is not Paid`, y 'Paid Invoices'. Las dos NO traen las mismas columnas
—la segunda no tiene `Rejected Reason` ni `Aging`— y `_merge_xlsx_files` rechaza
un header distinto al del primer chunk, que es justo lo que tiene que hacer.

`_normalise_part_headers` los reconcilia antes de que nadie los lea. Lo que se
fija aca:

1. que empareje por NOMBRE de columna y no por posicion (los dos tabs no traen
   las columnas en el mismo orden);
2. que una columna que un tab no tiene quede vacia, y ninguna fila se pierda;
3. que una columna que el primer chunk NO tiene frene la corrida en vez de
   descartarse en silencio — ese es el caso "el reporte cambio de nuevo";
4. que despues de normalizar, el merge y el invoice split trabajen sobre todo.
"""
import datetime as dt
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.scraper import (
    _merge_xlsx_files,
    _normalise_part_headers,
    _read,
    _require_non_empty_tab,
)

FULL = [
    "Urgency", "Entry ID", "PA Number", "Invoice #", "Client Name",
    "Client DDDID", "Client Number", "Service Code", "Status",
    "Rejected Reason", "Date Of Service", "Entry Creation Date", "Amount",
    "Aging",
]
# Como sale el tab de pagados: sin 'Rejected Reason' ni 'Aging'.
PAID = [c for c in FULL if c not in ("Rejected Reason", "Aging")]


def _row(header, entry_id, status, amount=10):
    values = {
        "Entry ID": entry_id,
        "Invoice #": f"INV{entry_id}",
        "Status": status,
        "Amount": amount,
        "Date Of Service": dt.datetime(2026, 1, 4),
        "Rejected Reason": "late" if status == "Rejected" else None,
        "Aging": "0-30",
    }
    return [values.get(name) for name in header]


def _make(path, header, rows, footer=True):
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    if footer:
        ws.append(["Applied filters: Date Of Service is on or after 01/01/2026"]
                  + [None] * (len(header) - 1))
    wb.save(path)
    return path


class NormaliseTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def test_a_narrower_part_is_widened_by_column_name(self):
        a = _make(self.d / "a.xlsx", FULL, [_row(FULL, "1", "Rejected")])
        b = _make(self.d / "b.xlsx", PAID, [_row(PAID, "2", "Paid")])
        _normalise_part_headers([a, b])

        rows = _read(b)
        self.assertEqual([str(c).strip() for c in rows[0]], FULL)
        row = rows[1]
        self.assertEqual(row[FULL.index("Entry ID")], "2")
        self.assertEqual(row[FULL.index("Status")], "Paid")
        self.assertEqual(row[FULL.index("Amount")], 10)
        # Las dos que el tab de pagados no trae quedan vacias, no corridas.
        self.assertIn(row[FULL.index("Rejected Reason")], (None, ""))
        self.assertIn(row[FULL.index("Aging")], (None, ""))

    def test_column_order_does_not_matter(self):
        """Se empareja por nombre: si el portal reordena, no se corren los datos."""
        shuffled = list(reversed(PAID))
        a = _make(self.d / "a.xlsx", FULL, [_row(FULL, "1", "Rejected")])
        b = _make(self.d / "b.xlsx", shuffled, [_row(shuffled, "9", "Paid", 77)])
        _normalise_part_headers([a, b])
        row = _read(b)[1]
        self.assertEqual(row[FULL.index("Entry ID")], "9")
        self.assertEqual(row[FULL.index("Amount")], 77)

    def test_no_row_is_lost(self):
        a = _make(self.d / "a.xlsx", FULL, [_row(FULL, "1", "Rejected")])
        b = _make(self.d / "b.xlsx", PAID,
                  [_row(PAID, str(n), "Paid") for n in range(2, 22)])
        _normalise_part_headers([a, b])
        self.assertEqual(len(_read(b)) - 1, 20)

    def test_an_unknown_column_stops_the_run(self):
        """Una columna que el primer chunk no tiene es un cambio de reporte, no
        algo para descartar callado."""
        a = _make(self.d / "a.xlsx", FULL, [_row(FULL, "1", "Rejected")])
        wider = FULL + ["Something New"]
        b = _make(self.d / "b.xlsx", wider, [_row(wider, "2", "Paid")])
        with self.assertRaises(ValueError) as ctx:
            _normalise_part_headers([a, b])
        self.assertIn("Something New", str(ctx.exception))

    def test_identical_headers_are_left_alone(self):
        a = _make(self.d / "a.xlsx", FULL, [_row(FULL, "1", "Rejected")])
        b = _make(self.d / "b.xlsx", FULL, [_row(FULL, "2", "Paid")])
        before = b.read_bytes()
        _normalise_part_headers([a, b])
        self.assertEqual(b.read_bytes(), before)


class MergeAfterNormaliseTests(unittest.TestCase):
    """El punto de todo esto: que las filas de los dos tabs terminen juntas."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def test_both_tabs_end_up_in_one_file(self):
        a = _make(self.d / "a.xlsx", FULL, [
            _row(FULL, "1", "Rejected"), _row(FULL, "2", "Pending"),
        ])
        b = _make(self.d / "b.xlsx", PAID, [
            _row(PAID, "3", "Paid"), _row(PAID, "4", "Paid"),
        ])
        _normalise_part_headers([a, b])
        out = _merge_xlsx_files([a, b], self.d / "out.xlsx")

        wb = load_workbook(out, read_only=True)
        try:
            rows = list(wb.active.iter_rows(values_only=True))
        finally:
            wb.close()  # Windows: sin esto el handle bloquea el tmpdir.
        self.assertEqual([str(c).strip() for c in rows[0]], FULL)
        ids = [r[FULL.index("Entry ID")] for r in rows[1:]]
        self.assertEqual(ids, ["1", "2", "3", "4"])

    def test_without_normalising_the_merge_still_refuses(self):
        """El guard de headers sigue vivo: normalizar es explicito, no implicito."""
        a = _make(self.d / "a.xlsx", FULL, [_row(FULL, "1", "Rejected")])
        b = _make(self.d / "b.xlsx", PAID, [_row(PAID, "2", "Paid")])
        with self.assertRaises(ValueError):
            _merge_xlsx_files([a, b], self.d / "out.xlsx")


class EmptyTabGuardTests(unittest.TestCase):
    """Un tab entero vacio tiene que frenar la corrida.

    Con un solo tab, un export vacio moria en el guard `total_rows == 0` de
    `_merge_xlsx_files`. Al pasar el invoice file a dos tabs ese guard dejo de
    poder verlo: suma los parts de los dos, y las filas de 'Paid Invoices'
    alcanzan para que el total nunca de cero aunque el tab que trae los
    rechazos vuelva vacio. El resultado seria un archivo solo-header entregado
    a ZipRide con el run en SUCCESS.
    """

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _meta(self, *counts):
        """`part_meta` tal como lo arma `_export_one_range`: path -> (ini, fin, filas)."""
        day = dt.date(2026, 9, 9)
        meta, paths = {}, []
        for i, rows in enumerate(counts):
            path = self.d / f"part_{i}.xlsx"
            meta[path] = (day, day, rows)
            paths.append(path)
        return paths, meta

    def test_the_merge_guard_alone_does_not_see_an_empty_tab(self):
        """El agujero que motiva el guard nuevo, escrito como test.

        Tab principal sin una sola fila, tab de pagados con filas: el merge
        produce un archivo y no levanta nada.
        """
        empty_main = _make(self.d / "main.xlsx", FULL, [])
        paid = _make(self.d / "paid.xlsx", FULL, [_row(FULL, "1", "Paid")])
        out = _merge_xlsx_files([empty_main, paid], self.d / "out.xlsx")
        self.assertTrue(out.exists())

    def test_a_tab_with_no_rows_stops_the_run(self):
        paths, meta = self._meta(0, 0, 0)
        with self.assertRaises(ValueError) as ctx:
            _require_non_empty_tab("Vendor Entry Status", paths, meta)
        self.assertIn("Vendor Entry Status", str(ctx.exception))
        self.assertIn("0 data rows", str(ctx.exception))

    def test_an_empty_chunk_inside_a_tab_is_still_legitimate(self):
        """El chunking adaptativo genera sub-rangos vacios; eso no es una falla."""
        paths, meta = self._meta(0, 1200, 0)
        _require_non_empty_tab("Paid Invoices", paths, meta)

    def test_a_tab_that_brought_rows_passes(self):
        paths, meta = self._meta(4795)
        _require_non_empty_tab("Vendor Entry Status", paths, meta)

    def test_a_path_without_meta_does_not_count_as_rows(self):
        """Un part sin entrada en `part_meta` no puede hacer pasar el guard."""
        paths, meta = self._meta(0)
        paths.append(self.d / "huerfano.xlsx")
        with self.assertRaises(ValueError):
            _require_non_empty_tab("Vendor Entry Status", paths, meta)


class DuplicateEntryIdTests(unittest.TestCase):
    """Un Entry ID que vuelve en los dos tabs se escribe UNA vez.

    Los dos tabs se exportan con minutos de diferencia contra datos vivos. Un
    invoice que el portal marca como pagado en esa ventana sale del tab de
    rechazos (exportado antes, cuando todavia figuraba sin pagar) y tambien del
    de pagados (exportado despues), con el mismo Entry ID. Sin deduplicar, el
    merge escribe las dos filas y el invoice split suma su Amount dos veces.

    Se conserva la ULTIMA: el tab de pagados se exporta al final, asi que es el
    dato mas fresco, y es la fila que `classify` ya habia elegido.

    Invariante medida sobre el archivo real del 2026-08-28: 100.462 filas,
    100.462 Entry IDs distintos, cero repetidos. Hasta que el archivo salio de
    dos tabs esto no podia pasar.
    """

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        # '2' cambia de estado entre los dos exports; '1' y '3' no se repiten.
        self.a = _make(self.d / "tab1.xlsx", FULL, [
            _row(FULL, "1", "Rejected", 100),
            _row(FULL, "2", "Rejected", 200),
        ])
        self.b = _make(self.d / "tab2.xlsx", FULL, [
            _row(FULL, "2", "Paid", 200),
            _row(FULL, "3", "Paid", 300),
        ])
        self.out = self.d / "merged.xlsx"

    def _merged_rows(self, keep=("1", "2", "3")):
        _merge_xlsx_files([self.a, self.b], self.out, frozenset(keep))
        wb = load_workbook(self.out, read_only=True)
        try:
            rows = list(wb.active.iter_rows(values_only=True))
        finally:
            wb.close()  # Windows: el handle bloquea el cleanup del tmpdir.
        return rows[0], rows[1:]

    def test_an_entry_in_both_tabs_is_written_once(self):
        _, rows = self._merged_rows()
        ids = [r[FULL.index("Entry ID")] for r in rows]
        self.assertEqual(sorted(ids), ["1", "2", "3"])

    def test_the_row_kept_is_the_one_from_the_later_tab(self):
        """La del tab de pagados: es el estado con el que el invoice quedo."""
        _, rows = self._merged_rows()
        by_id = {r[FULL.index("Entry ID")]: r for r in rows}
        self.assertEqual(by_id["2"][FULL.index("Status")], "Paid")

    def test_the_amount_is_not_counted_twice(self):
        _, rows = self._merged_rows()
        total = sum(r[FULL.index("Amount")] for r in rows)
        self.assertEqual(total, 600)  # 100 + 200 + 300, no 800

    def test_entries_that_do_not_repeat_are_all_kept(self):
        """El dedup no puede comerse filas distintas que comparten nada."""
        _, rows = self._merged_rows(keep=("1", "3"))
        ids = sorted(r[FULL.index("Entry ID")] for r in rows)
        self.assertEqual(ids, ["1", "3"])


if __name__ == "__main__":
    unittest.main()
