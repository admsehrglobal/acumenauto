"""QA of entries Paul adds while a run is already going (no browser, no database,
no email).

The list is read when the run starts and the payable pile goes out 15 to 20
minutes later. On 2026-09-17 Paul added 66 entries three minutes after a run had
read the list; the file ZipRide imported still carried all 66, and to him the
list was "not working". What is pinned here: whatever is on the list when a pile
is emailed comes out of it, and nothing else about the file changes.

The real `handle()` runs, with the real merge writing real files; only the model
layer, the browser and the email are stand-ins.
"""
import datetime as dt
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import django  # noqa: E402

django.setup()

from openpyxl import load_workbook  # noqa: E402

from app.invoice_split import PILE_PAYABLE, PILE_REJECTED  # noqa: E402
from app.management.commands import download_report as cmd  # noqa: E402
from app.scraper import _merge_xlsx_files  # noqa: E402
from tests.test_exception_drop import R1_HEADER, _make_xlsx, _r1  # noqa: E402

BURKE = ("NJ00006730", "Burke, M.")
BURKERT = ("NJ00006516", "Burkert, H.")


class _FakeRun:
    pk = 1

    def __init__(self):
        self.started_at = dt.datetime(2026, 9, 17, 20, 3, tzinfo=dt.timezone.utc)
        self.status = None
        self.error_message = ""
        self.filenames = ""
        self.exceptions_summary = ""
        self.finished_at = None

    def save(self):
        pass


def _rows(path):
    wb = load_workbook(path, read_only=True)
    try:
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()  # Windows: without this the handle blocks the tmpdir cleanup.


class LateEntryTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        self.run = _FakeRun()
        self.sent = {}  # pile -> rows of the file as it was emailed
        self.sent_bytes = {}
        self.written_bytes = {}
        self.stamped = None
        # The live list, as (key_1, key_2) rows per report; tests change it mid-run.
        self.live = {"invoices": [("100", "")], "auths": [], "accruals": []}
        self.during_download = None
        self.during_gap = None

        def _fake_send(items, recipients, label, override=None):
            for path, display_name in items:
                pile = PILE_REJECTED if PILE_REJECTED in display_name else PILE_PAYABLE
                self.sent[pile] = _rows(path)
                self.sent_bytes[pile] = path.read_bytes()
            return [(display_name, "msg-id") for _, display_name in items]

        def _fake_sleep(seconds):
            if self.during_gap:
                self.during_gap()

        def _capture_stamp(exceptions, complete=True):
            self.stamped = dict(exceptions)

        def _filter(**kwargs):
            query = mock.Mock()
            if "every_client_at__isnull" in kwargs:
                # No entry here is meant for every client.
                query.values_list.side_effect = lambda *a, **kw: []
            else:
                query.values_list.side_effect = lambda *a: list(self.live[kwargs["report"]])
            return query

        config = mock.Mock(report_1_enabled=True, report_2_enabled=False,
                           report_3_enabled=False, date_range_chunks=4)
        config.effective_dci_credentials.return_value = ("u", "p")
        recipients = mock.Mock()
        recipients.filter.return_value.values_list.return_value = ["paul@example.com"]
        run_mgr = mock.Mock()
        run_mgr.create.return_value = self.run

        for target, value in [
            ("send_reports_email", _fake_send),
            ("verify_delivery", lambda accepted: []),
            ("_notify_failure", lambda run: None),
            ("_stamp_matches_safely", _capture_stamp),
            ("Run", mock.Mock(objects=run_mgr, Status=cmd.Run.Status)),
            ("AppConfig", mock.Mock(load=lambda: config)),
            ("Recipient", mock.Mock(objects=recipients)),
            ("FileException", mock.Mock(objects=mock.Mock(filter=_filter))),
        ]:
            patcher = mock.patch.object(cmd, target, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        patcher = mock.patch.object(cmd.time, "sleep", _fake_sleep)
        patcher.start()
        self.addCleanup(patcher.stop)

    def _run(self, rejected_rows, payable_rows, keep=None):
        """`keep` maps a pile to the Entry IDs it keeps, as the invoice split
        does; an empty set gives the header-only pile of a day with none."""
        async def _fake_download(**kwargs):
            # What the scraper does: both piles merged with the list read at start.
            spec = kwargs["chunked_reports"][0].exceptions
            items = []
            for pile, rows in ((PILE_REJECTED, rejected_rows), (PILE_PAYABLE, payable_rows)):
                chunk = self.d / f"chunk_{pile}.xlsx"
                _make_xlsx(chunk, R1_HEADER, rows)
                out = self.d / f"{pile}.xlsx"
                _merge_xlsx_files([chunk], out, (keep or {}).get(pile), spec)
                self.written_bytes[pile] = out.read_bytes()
                items.append((out, f"R1 - {pile} (2025-06-08 to 2026-09-17)"))
            if self.during_download:
                self.during_download()
            for item in items:
                kwargs["on_report_ready"](*item)
            return items

        with mock.patch.object(cmd, "download_reports", _fake_download):
            cmd.Command().handle(output_dir=str(self.d), no_email=False, reports="1")

    def _invoices(self, pile):
        column = R1_HEADER.index("Invoice #")
        return sorted(str(r[column]) for r in self.sent[pile][1:])

    def test_an_entry_added_while_the_run_downloads_comes_out_of_both_piles(self):
        self.during_download = lambda: self.live["invoices"].append(("200", ""))
        self._run(
            [_r1("1", "100", "Rejected"), _r1("2", "200", "Rejected"), _r1("3", "300", "Rejected")],
            [_r1("4", "100"), _r1("5", "200"), _r1("6", "400")],
        )
        self.assertEqual(self._invoices(PILE_REJECTED), ["300"])
        self.assertEqual(self._invoices(PILE_PAYABLE), ["400"])

    def test_an_entry_added_during_the_gap_comes_out_of_the_payable_pile(self):
        """The rejections are already in ZipRide by then; the payable pile is not."""
        self.during_gap = lambda: self.live["invoices"].append(("200", ""))
        self._run(
            [_r1("1", "200", "Rejected")],
            [_r1("2", "200"), _r1("3", "400")],
        )
        self.assertEqual(self._invoices(PILE_REJECTED), ["200"])
        self.assertEqual(self._invoices(PILE_PAYABLE), ["400"])

    def test_a_client_named_entry_added_mid_run_takes_only_that_clients_line(self):
        """What the one-click fix on the page writes, clicked while a run is going."""
        self.during_download = lambda: self.live["invoices"].append(("119344", BURKE[0]))
        self._run(
            [_r1("9", "300", "Rejected")],
            [_r1("1", "119344", client=BURKE[0], name=BURKE[1]),
             _r1("2", "119344", client=BURKERT[0], name=BURKERT[1])],
        )
        client = R1_HEADER.index("Client Number")
        self.assertEqual([r[client] for r in self.sent[PILE_PAYABLE][1:]], [BURKERT[0]])

    def test_with_nothing_new_the_file_goes_out_exactly_as_the_merge_wrote_it(self):
        self._run([_r1("1", "300", "Rejected")], [_r1("2", "100"), _r1("3", "400")])
        for pile in (PILE_REJECTED, PILE_PAYABLE):
            self.assertEqual(self.sent_bytes[pile], self.written_bytes[pile], pile)

    def test_the_rewritten_file_keeps_every_column_and_value_of_the_rows_it_keeps(self):
        """The columns are a contract with Juan's import; only rows may go."""
        kept = _r1("3", "400")
        self.during_download = lambda: self.live["invoices"].append(("200", ""))
        self._run([_r1("9", "300", "Rejected")], [_r1("2", "200"), kept])
        header, *rows = self.sent[PILE_PAYABLE]
        self.assertEqual(header, R1_HEADER)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][R1_HEADER.index("Entry ID")], "3")
        self.assertEqual(rows[0][R1_HEADER.index("Date Of Service")], kept[10])
        self.assertEqual(rows[0][R1_HEADER.index("Amount")], kept[12])

    def test_the_run_record_counts_the_late_entry(self):
        self.during_download = lambda: self.live["invoices"].append(("200", ""))
        self._run(
            [_r1("1", "200", "Rejected"), _r1("4", "300", "Rejected")],
            [_r1("2", "100"), _r1("3", "200"), _r1("5", "400")],
        )
        self.assertEqual(
            self.run.exceptions_summary,
            "Invoices: 3 rows dropped (2 of 2 keys matched)",
        )

    def test_every_row_a_late_entry_takes_is_counted_and_only_what_matched(self):
        self.during_download = lambda: self.live["invoices"].extend([("200", ""), ("999", "")])
        self._run(
            [_r1("1", "300", "Rejected")],
            [_r1("2", "100"), _r1("3", "200"), _r1("4", "200"), _r1("5", "400")],
        )
        self.assertEqual(
            self.run.exceptions_summary,
            "Invoices: 3 rows dropped (2 of 3 keys matched)",
        )

    def test_an_entry_narrowed_before_its_pile_went_out_is_not_reported_as_shared(self):
        """Added with no client during the download, then fixed with one click in
        the gap: it took nothing, so the run must not say it took two clients."""
        def narrow():
            self.live["invoices"].remove(("119344", ""))
            self.live["invoices"].append(("119344", BURKE[0]))

        self.during_download = lambda: self.live["invoices"].append(("119344", ""))
        self.during_gap = narrow
        self._run(
            [_r1("9", "300", "Rejected")],
            [_r1("1", "119344", client=BURKE[0], name=BURKE[1]),
             _r1("2", "119344", client=BURKERT[0], name=BURKERT[1]),
             _r1("3", "400")],
        )
        client = R1_HEADER.index("Client Number")
        self.assertEqual(
            sorted(r[client] for r in self.sent[PILE_PAYABLE][1:]), ["C1", BURKERT[0]]
        )
        self.assertNotIn("with no Client Number", self.run.exceptions_summary)

    def test_a_pile_with_no_rows_and_a_late_entry_logs_no_error(self):
        """A day with no rejections: the pile is a header only, and still goes out."""
        self.during_download = lambda: self.live["invoices"].append(("200", ""))
        with self.assertNoLogs(cmd.logger, level="ERROR"):
            self._run(
                [_r1("1", "300", "Rejected")],
                [_r1("2", "200"), _r1("3", "400")],
                keep={PILE_REJECTED: set()},
            )
        self.assertEqual(self.sent[PILE_REJECTED], [R1_HEADER])
        self.assertEqual(self._invoices(PILE_PAYABLE), ["400"])

    def test_a_pile_the_start_list_emptied_and_a_late_entry_logs_no_error(self):
        self.during_download = lambda: self.live["invoices"].append(("200", ""))
        with self.assertNoLogs(cmd.logger, level="ERROR"):
            self._run([_r1("1", "100", "Rejected")], [_r1("2", "200"), _r1("3", "400")])
        self.assertNotIn(PILE_REJECTED, self.sent)
        self.assertEqual(self._invoices(PILE_PAYABLE), ["400"])

    def test_a_file_the_invoice_merge_did_not_write_is_left_alone(self):
        """The auth file goes through the same `_send`; the invoice list is not its list."""
        other = self.d / "auths.xlsx"
        _make_xlsx(other, R1_HEADER, [_r1("1", "200")])
        before = other.read_bytes()
        self.live["invoices"].append(("200", ""))
        spec = cmd.make_drop_spec("invoices", [("100", "")])
        cmd._drop_late_entries(other, {"invoices": spec}, spec.keys)
        self.assertEqual(other.read_bytes(), before)

    def test_the_late_entry_is_stamped_as_matched(self):
        self.during_download = lambda: self.live["invoices"].append(("200", ""))
        self._run([_r1("9", "300", "Rejected")], [_r1("2", "200"), _r1("3", "400")])
        spec = self.stamped["invoices"]
        matched = set().union(*(keys for _, keys in spec.stats.values()))
        self.assertIn(("200", ""), matched)

    def test_a_late_entry_that_empties_the_payable_pile_holds_it_back(self):
        self.during_gap = lambda: self.live["invoices"].append(("200", ""))
        self._run([_r1("1", "300", "Rejected")], [_r1("2", "200")])
        self.assertNotIn(PILE_PAYABLE, self.sent)
        self.assertIn("1 file(s) not emailed (no rows left)", self.run.exceptions_summary)

    def test_if_reading_the_list_again_fails_the_file_still_goes_out(self):
        self.during_download = lambda: self.live["invoices"].append(("200", ""))
        real = cmd._load_list
        calls = []

        def _flaky(slug):
            calls.append(slug)
            if len(calls) > 3:  # the three reads at start work, the ones after fail
                raise RuntimeError("database went away")
            return real(slug)

        with mock.patch.object(cmd, "_load_list", _flaky):
            self._run([_r1("1", "300", "Rejected")], [_r1("2", "200"), _r1("3", "400")])
        self.assertEqual(self._invoices(PILE_PAYABLE), ["200", "400"])
        self.assertEqual(self.run.status, cmd.Run.Status.SUCCESS)


if __name__ == "__main__":
    unittest.main()
