"""QA of applying the File Exceptions lists while the files are written (no browser).

Covers what would be silent if it broke:

1. `_merge_xlsx_files` with a `DropSpec` leaves out every row of an excepted
   key and nothing else, resolving the key columns by name, and records what it
   dropped per output file. A key typed as a number matches the text cell.
2. The composite keys of the auths and accruals files need BOTH parts to match:
   five PA Numbers sit under two different clients in the real file.
2b. The invoices key takes the client too, and it is OPTIONAL: an entry with no
   client drops every row carrying the number (which is every entry stored
   before 2026-09-21) and one with a client drops only that client's row.
3. `_split_for_email` applies the same list to every file it re-merges, and the
   big merge it deletes is not counted twice.
4. R2 is rewritten in place: same path, same sheet name, same header, the
   excepted rows gone.
"""
import datetime as dt
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from app import scraper
from app.file_exceptions import make_drop_spec, summarize
from app.scraper import (
    ChunkedReport,
    _apply_exceptions_in_place,
    _merge_xlsx_files,
    _split_for_email,
)

R1_HEADER = [
    "Urgency", "Entry ID", "PA Number", "Invoice #", "Client Name", "Client DDDID",
    "Client Number", "Service Code", "Status", "Rejected Reason", "Date Of Service",
    "Entry Creation Date", "Amount", "Aging",
]
R3_HEADER = [
    "Client Name", "Client DDDID", "PA Number", "Start Date",
    "End Date", "Accrual Schedule Date", "Accrual Schedule Amount",
]
R2_HEADER = ["Authorization ID", "Client Name", "Client ID", "Client DDDID", "Status"]


def _r1(entry_id, invoice, status="Paid", client="C1", name="Cli"):
    return ["", entry_id, "PA1", invoice, name, "D1", client, "Transportation",
            status, "", dt.datetime(2025, 7, 20), dt.datetime(2025, 7, 28),
            84.0, 0]


def _r3(dddid, pa, amount=100):
    return ["Cli", dddid, pa, dt.datetime(2025, 6, 8), dt.datetime(2027, 1, 1),
            dt.datetime(2026, 1, 5), amount]


def _make_xlsx(path, header, rows, title=None):
    wb = Workbook()
    ws = wb.active
    if title:
        ws.title = title
    ws.append(header)
    for r in rows:
        ws.append(r)
    # PBI appends this row to every export; the merge skips it.
    ws.append(["Applied filters: EndDate is on or after X"] + [None] * (len(header) - 1))
    wb.save(path)


def _read(path):
    wb = load_workbook(path, read_only=True)
    try:
        return wb.sheetnames[0], list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()  # Windows: without this the handle blocks the tmpdir cleanup.


class InvoiceDropTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        self.chunk = self.d / "chunk.xlsx"
        _make_xlsx(self.chunk, R1_HEADER, [
            _r1("1", "500", "Rejected"),
            _r1("2", "500"),
            _r1("3", "600"),
            _r1("4", "700"),
        ])

    def _entry_ids(self, path):
        return [r[1] for r in _read(path)[1][1:]]

    def test_every_row_of_an_excepted_invoice_is_left_out_of_the_pile(self):
        drop = make_drop_spec("invoices", [("500", "")])
        out = self.d / "payable.xlsx"
        _merge_xlsx_files([self.chunk], out, frozenset({"2", "3", "4"}), drop)
        self.assertEqual(self._entry_ids(out), ["3", "4"])
        # Counted per output file: row 1 was never in this pile to begin with.
        self.assertEqual(drop.stats[out], (1, {("500", "")}))

    def test_a_number_in_the_list_matches_the_text_cell_in_the_export(self):
        """Paul's spreadsheet will carry 600, the export carries '600'."""
        drop = make_drop_spec("invoices", [(600.0, "")])
        out = self.d / "all.xlsx"
        _merge_xlsx_files([self.chunk], out, None, drop)
        self.assertEqual(self._entry_ids(out), ["1", "2", "4"])

    def test_no_match_writes_everything_and_says_so(self):
        drop = make_drop_spec("invoices", [("999", "")])
        out = self.d / "all.xlsx"
        _merge_xlsx_files([self.chunk], out, None, drop)
        self.assertEqual(self._entry_ids(out), ["1", "2", "3", "4"])
        self.assertEqual(drop.stats[out], (0, set()))

    def test_a_missing_key_column_fails_loudly(self):
        """Same policy as the invoice split: a renamed column must not turn into
        a file that silently still carries the excepted rows."""
        header = [c for c in R1_HEADER if c != "Invoice #"]
        row = _r1("1", "500")
        del row[R1_HEADER.index("Invoice #")]
        chunk = self.d / "drifted.xlsx"
        _make_xlsx(chunk, header, [row])
        with self.assertRaises(ValueError) as ctx:
            _merge_xlsx_files([chunk], self.d / "out.xlsx", None,
                              make_drop_spec("invoices", [("1", "")]))
        self.assertIn("Invoice #", str(ctx.exception))


class AccrualDropTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        self.chunk = self.d / "chunk.xlsx"
        _make_xlsx(self.chunk, R3_HEADER, [
            _r3("431798", "1553308787"),
            _r3("653983", "1553308787"),
            [None] * 7,  # the blank row PBI leaves before its footer
            _r3(None, "1553411994"),  # the 'VOID, V.' rows: no client id
            _r3("306194", "1553411994"),
        ])

    def _keys(self, path):
        return [(r[1], r[2]) for r in _read(path)[1][1:]]

    def test_both_parts_must_match(self):
        """Solo sale la fila cuya clave completa esta en la lista.

        La fila en blanco que PBI deja antes del footer ya no viaja al archivo:
        desde el 2026-09-05 no cuenta como dato, que es lo que mantenia
        desactivado el guard de "reporte vacio" (cuatro chunks sin una sola fila
        daban total_rows == 4 y salia un archivo con solo el header).
        """
        drop = make_drop_spec("accruals", [("431798", "1553308787")])
        out = self.d / "out.xlsx"
        _merge_xlsx_files([self.chunk], out, None, drop)
        self.assertEqual(self._keys(out), [
            ("653983", "1553308787"),
            (None, "1553411994"),
            ("306194", "1553411994"),
        ])

    def test_rows_without_a_client_id_are_never_matched(self):
        """Neither the blank row nor a PA with an empty DDDID can be an
        exception: a list entry always has both parts."""
        drop = make_drop_spec("accruals", [("", "1553411994")])
        self.assertIsNone(drop)


class SplitDropTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        self.parts = []
        self.part_meta = {}
        d0 = dt.date(2025, 1, 1)
        for i in range(4):
            p = self.d / f"r3_part_{i:03d}.xlsx"
            _make_xlsx(p, R3_HEADER, [_r3("1", f"P{i}-{j}") for j in range(25)])
            start = d0 + dt.timedelta(days=10 * i)
            self.part_meta[p] = (start, start + dt.timedelta(days=9), 25)
            self.parts.append(p)
        self.report_range = (d0, d0 + dt.timedelta(days=39))
        self.merged = self.d / "r3_merged.xlsx"

    def test_every_output_file_honours_the_list_and_nothing_is_counted_twice(self):
        drop = make_drop_spec("accruals", [("1", "P0-3"), ("1", "P2-7"), ("1", "P3-24")])
        _merge_xlsx_files(self.parts, self.merged, None, drop)
        size = self.merged.stat().st_size

        with patch.object(scraper, "_ATTACHMENT_MAX_BYTES", size // 2), \
                patch.object(scraper, "_ATTACHMENT_TARGET_BYTES", size // 4):
            outputs = _split_for_email(
                self.merged, self.parts, self.part_meta, self.report_range,
                self.d, "r3", "ts", None, drop,
            )

        self.assertGreater(len(outputs), 1)
        pas = [r[2] for path, _, _ in outputs for r in _read(path)[1][1:]]
        self.assertEqual(len(pas), 97)
        self.assertNotIn("P0-3", pas)
        self.assertNotIn("P2-7", pas)
        self.assertNotIn("P3-24", pas)
        # The deleted merge is gone from the record; the parts add up to 3.
        self.assertEqual(set(drop.stats), {path for path, _, _ in outputs})
        self.assertEqual(sum(n for n, _ in drop.stats.values()), 3)


class RawExportRewriteTests(unittest.TestCase):
    """R2 is the raw Power BI download; the list is applied by rewriting it."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        self.raw = self.d / "view_vendor_authorization_report_ts.xlsx"
        _make_xlsx(self.raw, R2_HEADER, [
            [173066812, "A", "NJ1", "721253", "Approved"],
            [173066813, "B", "NJ2", "721253", "Approved"],
            [173066814, "C", "NJ3", "721254", "Approved"],
        ], title="Export")

    def test_rewritten_in_place_without_the_excepted_rows(self):
        drop = make_drop_spec("auths", [("721253", "173066813"), ("721254", "173066812")])
        _apply_exceptions_in_place(self.raw, drop)

        self.assertTrue(self.raw.exists())
        self.assertEqual(list(self.d.iterdir()), [self.raw])
        sheet, rows = _read(self.raw)
        self.assertEqual(sheet, "Export")
        self.assertEqual(list(rows[0]), R2_HEADER)
        self.assertEqual([r[0] for r in rows[1:]], [173066812, 173066814])
        self.assertEqual(sum(n for n, _ in drop.stats.values()), 1)


class EmptiedByTheListTests(unittest.TestCase):
    """A file the list left with no data rows is remembered, so the command can
    skip that email (Juan Pablo, 2026-09-07). Only that case: a file that had
    nothing to begin with is not marked, because that one still goes out."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def test_a_merge_the_list_emptied_is_marked_under_the_output_path(self):
        part = self.d / "part.xlsx"
        _make_xlsx(part, R1_HEADER, [_r1(1, "500"), _r1(2, "500")])
        drop = make_drop_spec("invoices", [("500",)])
        out = self.d / "merged.xlsx"
        _merge_xlsx_files([part], out, None, drop)
        self.assertEqual(drop.emptied, {out})

    def test_a_merge_with_survivors_is_not_marked(self):
        part = self.d / "part.xlsx"
        _make_xlsx(part, R1_HEADER, [_r1(1, "500"), _r1(2, "501")])
        drop = make_drop_spec("invoices", [("500",)])
        out = self.d / "merged.xlsx"
        _merge_xlsx_files([part], out, None, drop)
        self.assertEqual(drop.emptied, set())

    def test_a_pile_that_was_already_empty_is_not_marked(self):
        """No rejected invoices today is an answer, not a gap: that file still
        goes out, so nothing may mark it."""
        part = self.d / "part.xlsx"
        _make_xlsx(part, R1_HEADER, [_r1(1, "500")])
        drop = make_drop_spec("invoices", [("999",)])
        out = self.d / "merged.xlsx"
        _merge_xlsx_files([part], out, frozenset(), drop)
        self.assertEqual(drop.emptied, set())

    def test_r2_is_marked_under_the_path_the_command_will_look_up(self):
        """The rewrite merges into a temporary name and then renames. The mark
        has to end up on the ORIGINAL path: that is the one `_send` is holding,
        and if it stays on the temporary name the skip never fires for R2."""
        raw = self.d / "view_vendor_authorization_report_ts.xlsx"
        _make_xlsx(raw, R2_HEADER, [
            [173066812, "A", "NJ1", "721253", "Approved"],
        ], title="Export")
        drop = make_drop_spec("auths", [("721253", "173066812")])
        _apply_exceptions_in_place(raw, drop)
        self.assertEqual(drop.emptied, {raw})
        self.assertEqual(list(self.d.iterdir()), [raw])

    def test_forgetting_the_big_merge_clears_its_mark_too(self):
        part = self.d / "part.xlsx"
        _make_xlsx(part, R1_HEADER, [_r1(1, "500")])
        drop = make_drop_spec("invoices", [("500",)])
        out = self.d / "merged.xlsx"
        _merge_xlsx_files([part], out, None, drop)
        self.assertEqual(drop.emptied, {out})
        drop.forget(out)
        self.assertEqual(drop.emptied, set())


class SpecTests(unittest.TestCase):
    def test_no_exceptions_unless_the_command_passes_them(self):
        spec = ChunkedReport(
            url="u", button_name="b", n_chunks=4, today=dt.date(2026, 9, 3),
            tab_name=None, single_slicer=True, full_range=True,
        )
        self.assertIsNone(spec.exceptions)


if __name__ == "__main__":
    unittest.main()


class InvoiceClientKeyTests(unittest.TestCase):
    """The optional client half of the invoices key.

    ZipRide's import of the 2026-09-17 file flagged 66 entries, 61 of them
    "Client mismatch: expected <one client>, got <another>". Measured on the
    2026-09-06 pair, dropping those 66 by number alone takes 79 rows, not 70:
    the 9 extra ones belong to the clients ZipRide says the invoice is for.
    Naming the client is what stops that, and leaving it empty has to keep
    behaving exactly as the list did before.
    """

    HEADER_NO_CLIENT = [c for c in R1_HEADER if c != "Client Number"]

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)
        self.chunk = self.d / "chunk.xlsx"
        # Invoice 163746 under two clients, which is the real shape: Paul asked
        # for the NJ00001544 row and the NJ00013618 one is another client's.
        _make_xlsx(self.chunk, R1_HEADER, [
            _r1("1", "163746", client="NJ00001544"),
            _r1("2", "163746", client="NJ00013618"),
            _r1("3", "999", client="NJ00001544"),
        ])

    def _ids(self, path):
        return [r[1] for r in _read(path)[1][1:]]

    def _merge(self, drop, chunk=None):
        out = self.d / "payable.xlsx"
        _merge_xlsx_files([chunk or self.chunk], out, None, drop)
        return out

    def test_a_client_narrows_the_entry_to_that_client_s_row(self):
        drop = make_drop_spec("invoices", [("163746", "NJ00001544")])
        out = self._merge(drop)
        self.assertEqual(self._ids(out), ["2", "3"])
        self.assertEqual(drop.stats[out], (1, {("163746", "nj00001544")}))

    def test_no_client_still_drops_every_row_with_the_number(self):
        """What the 428 entries already on the list do, unchanged."""
        drop = make_drop_spec("invoices", [("163746", "")])
        out = self._merge(drop)
        self.assertEqual(self._ids(out), ["3"])
        self.assertEqual(drop.stats[out], (2, {("163746", "")}))

    def test_the_client_is_matched_ignoring_case(self):
        drop = make_drop_spec("invoices", [("163746", "nj00001544")])
        self.assertEqual(self._ids(self._merge(drop)), ["2", "3"])

    def test_both_shapes_of_one_number_are_each_recorded_as_matched(self):
        """A row is dropped once, but the page must not call either entry a
        typo: 'Checked, matched nothing' is how a wrong key is spotted."""
        drop = make_drop_spec(
            "invoices", [("163746", ""), ("163746", "NJ00001544")]
        )
        out = self._merge(drop)
        self.assertEqual(self._ids(out), ["3"])
        dropped, matched = drop.stats[out]
        self.assertEqual(dropped, 2)
        self.assertEqual(matched, {("163746", ""), ("163746", "nj00001544")})

    def test_an_export_without_the_client_column_does_not_fail_the_run(self):
        """Acumen has renamed R1 columns twice in 2026. A rename of the column
        the key does not require must not kill the invoice file; the entries
        that name a client simply stop matching."""
        chunk = self.d / "no_client.xlsx"
        _make_xlsx(chunk, self.HEADER_NO_CLIENT, [
            [c for i, c in enumerate(_r1("1", "163746", client="NJ00001544"))
             if R1_HEADER[i] != "Client Number"],
            [c for i, c in enumerate(_r1("3", "999"))
             if R1_HEADER[i] != "Client Number"],
        ])
        drop = make_drop_spec(
            "invoices", [("163746", "NJ00001544"), ("999", "")]
        )
        out = self._merge(drop, chunk)
        # The wildcard still works; the one naming a client matched nothing.
        self.assertEqual(self._ids(out), ["1"])
        self.assertEqual(drop.stats[out], (1, {("999", "")}))

    def test_a_missing_required_column_still_fails_loudly(self):
        chunk = self.d / "no_invoice.xlsx"
        header = [c for c in R1_HEADER if c != "Invoice #"]
        _make_xlsx(chunk, header, [
            [c for i, c in enumerate(_r1("1", "163746"))
             if R1_HEADER[i] != "Invoice #"],
        ])
        drop = make_drop_spec("invoices", [("163746", "")])
        with self.assertRaises(ValueError):
            self._merge(drop, chunk)


class SharedNumberTests(unittest.TestCase):
    """Which invoice numbers the written files carry on more than one client's
    lines, and which of those are on the list with no client.

    Paul's list of 2026-09-17 carried 66 numbers with no client. Nine of them
    were also on another client's line, one ZipRide was accepting, and those
    nine lines ($1,914.50) were left out of every file for a week without
    anything saying so. On 2026-09-24 five more entries on the live list had
    the same shape (1, 2, 26, 32 and 119344).
    """

    BURKE = (("NJ00006730",), "Burke, M.")
    BURKERT = (("NJ00006516",), "Burkert, H.")

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.d = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def _merge(self, rows, drop, keep=None, name="payable.xlsx", header=R1_HEADER):
        chunk = self.d / f"chunk_{name}"
        _make_xlsx(chunk, header, rows)
        out = self.d / name
        _merge_xlsx_files([chunk], out, keep, drop)
        return out

    def _shared_rows(self):
        return [
            _r1("1", "119344", client="NJ00006730", name="Burke, M."),
            _r1("2", "119344", client="NJ00006516", name="Burkert, H."),
            _r1("3", "999", client="NJ00006730", name="Burke, M."),
        ]

    def test_a_listed_number_on_two_clients_lines_is_named_with_both(self):
        drop = make_drop_spec("invoices", [("119344", ""), ("999", "")])
        self._merge(self._shared_rows(), drop)
        both = [self.BURKERT, self.BURKE]
        self.assertEqual(drop.shared(), {("119344",): both})
        self.assertEqual(drop.shared_entries(), {("119344",): both})
        self.assertIn(
            ", 1 with no Client Number matched more than one Client Number",
            summarize([drop]),
        )

    def test_every_shared_number_is_known_whether_listed_or_not(self):
        """So the page can flag one the moment Paul adds it, not a run later."""
        drop = make_drop_spec("invoices", [("999", "")])
        self._merge(self._shared_rows(), drop)
        self.assertEqual(list(drop.shared()), [("119344",)])
        self.assertEqual(drop.shared_entries(), {})
        self.assertNotIn("more than one", summarize([drop]))

    def test_an_entry_naming_its_client_is_not_a_shared_entry(self):
        drop = make_drop_spec("invoices", [("119344", "NJ00006730")])
        self._merge(self._shared_rows(), drop)
        self.assertEqual(drop.shared_entries(), {})

    def test_a_number_on_one_client_s_lines_is_not_shared(self):
        """Two weeks of service under one invoice is the common shape (612
        invoice-and-client pairs on the 2026-09-06 file), not a shared number."""
        drop = make_drop_spec("invoices", [("500", "")])
        self._merge([
            _r1("1", "500", client="NJ00001544"),
            _r1("2", "500", client="NJ00001544"),
        ], drop)
        self.assertEqual(drop.shared(), {})

    def test_one_client_in_two_spellings_is_one_client(self):
        drop = make_drop_spec("invoices", [("500", "")])
        self._merge([
            _r1("1", "500", client="NJ00001544"),
            _r1("2", "500", client="nj00001544"),
        ], drop)
        self.assertEqual(drop.shared(), {})

    def test_a_row_with_no_client_does_not_count_as_a_second_client(self):
        drop = make_drop_spec("invoices", [("500", "")])
        self._merge([
            _r1("1", "500", client="NJ00001544"),
            _r1("2", "500", client=""),
        ], drop)
        self.assertEqual(drop.shared(), {})

    def test_the_two_clients_can_be_in_the_two_files(self):
        """The invoice file goes out as two files, rejected and payable, each
        written by its own merge. One client in each is still two clients."""
        drop = make_drop_spec("invoices", [("119344", "")])
        rows = [
            _r1("1", "119344", "Rejected", client="NJ00006730"),
            _r1("2", "119344", client="NJ00006516"),
        ]
        self._merge(rows, drop, keep=frozenset({"1"}), name="rejected.xlsx")
        self._merge(rows, drop, keep=frozenset({"2"}), name="payable.xlsx")
        self.assertEqual(
            drop.shared_entries(),
            {("119344",): [(("NJ00006516",), "Cli"), (("NJ00006730",), "Cli")]},
        )

    def test_forgetting_a_merge_keeps_what_it_saw(self):
        """`_split_for_email` forgets the big merge it deletes, but the files
        that replace it carry the same rows: nothing to take back."""
        drop = make_drop_spec("invoices", [("119344", "")])
        out = self._merge(self._shared_rows(), drop)
        drop.forget(out)
        self.assertEqual(list(drop.shared_entries()), [("119344",)])

    def test_without_the_client_column_nothing_is_shared(self):
        """Rows with no client cannot be told apart, so the run stays quiet
        rather than guess."""
        header = [c for c in R1_HEADER if c != "Client Number"]
        rows = [
            [c for i, c in enumerate(r) if R1_HEADER[i] != "Client Number"]
            for r in self._shared_rows()
        ]
        drop = make_drop_spec("invoices", [("119344", "")])
        self._merge(rows, drop, header=header)
        self.assertEqual(drop.shared(), {})

    def test_without_the_name_column_the_client_number_is_enough(self):
        header = [c for c in R1_HEADER if c != "Client Name"]
        rows = [
            [c for i, c in enumerate(r) if R1_HEADER[i] != "Client Name"]
            for r in self._shared_rows()
        ]
        drop = make_drop_spec("invoices", [("119344", "")])
        self._merge(rows, drop, header=header)
        self.assertEqual(
            drop.shared(),
            {("119344",): [(("NJ00006516",), ""), (("NJ00006730",), "")]},
        )

    def test_a_key_with_no_optional_part_tracks_nothing(self):
        drop = make_drop_spec("auths", [("721253", "173066812")])
        self.assertIsNone(drop.owners)
        self.assertEqual(drop.shared(), {})
